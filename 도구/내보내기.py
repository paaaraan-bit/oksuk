"""원고/ 폴더의 장 파일을 하나로 합쳐 내보낸다.

- 내보내기/전체원고.txt : 한 문단을 한 줄로 되돌린 텍스트
- 내보내기/전체원고.docx : python-docx가 설치되어 있으면 함께 만든다

사용법: python3 도구/내보내기.py
"""
import pathlib

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "원고"
OUT = ROOT / "내보내기"
BOOK_TITLE = "저승 알바천재 노수오"
SCENE_BREAK = "* * *"


def read_chapter(path):
    """장 파일을 (제목, 문단 목록)으로 읽는다. 장면 전환은 SCENE_BREAK 문단으로 남긴다."""
    lines = path.read_text(encoding="utf-8").split("\n")
    title = lines[0].lstrip("#").strip()
    paras, buf = [], []
    for line in lines[1:] + [""]:
        if line.strip():
            buf.append(line.strip())
        elif buf:
            paras.append(" ".join(buf))
            buf = []
    return title, paras


def main():
    chapters = [read_chapter(p) for p in sorted(SRC.glob("*.md"))]
    OUT.mkdir(exist_ok=True)

    txt = [BOOK_TITLE, ""]
    for title, paras in chapters:
        txt += ["", title, ""]
        txt += ["" if p == SCENE_BREAK else p for p in paras]
    (OUT / "전체원고.txt").write_text("\n".join(txt) + "\n", encoding="utf-8")
    print("만듦:", OUT / "전체원고.txt")

    try:
        import docx
    except ImportError:
        print("워드 파일은 건너뜀 (pip install python-docx 후 다시 실행)")
        return
    doc = docx.Document()
    doc.add_heading(BOOK_TITLE, level=0)
    for title, paras in chapters:
        doc.add_heading(title, level=1)
        for p in paras:
            doc.add_paragraph("" if p == SCENE_BREAK else p)
    doc.save(OUT / "전체원고.docx")
    print("만듦:", OUT / "전체원고.docx")


if __name__ == "__main__":
    main()
